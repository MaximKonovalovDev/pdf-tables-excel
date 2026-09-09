"""
PDF Tables to Excel - pull the tables out of PDF files into .xlsx sheets.

- Runs 100% on your PC. No cloud, no account, no AI calls.
- Two read modes plus an optional OCR fallback:
    * lattice : reads tables drawn with ruling lines (most accurate)
    * stream  : guesses columns from text spacing (no lines needed)
    * ocr     : last resort for scanned pages (needs free Tesseract)
- Every table gets an honest confidence label: HIGH / MEDIUM / LOW.
- Writes a preview.html so you can eyeball every table before trusting it.

Use it two ways:

    python tables.py                 -> opens the desktop window
    python tables.py file.pdf        -> writes file.xlsx + file.preview.html
    python tables.py folder --out-dir done --ocr   -> batch mode

CPU-only. Python 3.10+. Windows/macOS/Linux.
"""

from __future__ import annotations

import argparse
import bisect
import html
import os
import queue
import statistics
import sys
import threading
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

APP = "PDF Tables to Excel"
VERSION = "1.0.0"

METHOD_LABELS = {
    "lattice": "Lined grid (exact)",
    "stream": "Spacing guess",
    "ocr": "Scanned-page OCR",
}

# Honest one-line explanations shown next to every table in the preview,
# the console log, and the Excel summary sheet.
NOTES = {
    ("lattice", "HIGH"): "Read straight from the PDF's own ruling lines.",
    ("lattice", "MEDIUM"): "From ruling lines; a few cells may need cleanup.",
    ("lattice", "LOW"): "Lines were faint - check this one carefully.",
    ("stream", "HIGH"): "Column layout inferred from text spacing on clean text.",
    ("stream", "MEDIUM"): "Layout guessed from text spacing - spot-check merged cells.",
    ("stream", "LOW"): "Rough guess from spacing only - verify before you trust it.",
    ("ocr", "LOW"): "Reconstructed from a scanned picture of the page. "
    "OCR makes mistakes - proofread every number before use.",
}


def label_for(accuracy, method):
    """Map a detection score to an honest HIGH/MEDIUM/LOW label."""
    if method == "ocr":
        return "LOW"  # never oversell OCR output
    if accuracy is None:
        return "LOW"
    if accuracy >= 90:
        return "HIGH"
    if accuracy >= 70:
        return "MEDIUM"
    return "LOW"


def note_for(method, label):
    return NOTES.get((method, label), "")


@dataclass
class TableResult:
    """One detected table. `rows` is a list of lists of plain strings."""

    page: int
    index: int          # 1-based table number within its page
    method: str         # lattice | stream | ocr
    accuracy: float | None
    rows: list

    @property
    def label(self) -> str:
        return label_for(self.accuracy, self.method)

    @property
    def note(self) -> str:
        return note_for(self.method, self.label)

    @property
    def size(self) -> str:
        cols = max((len(r) for r in self.rows), default=0)
        return f"{len(self.rows)} rows x {cols} cols"

    @property
    def sheet_name(self) -> str:
        return f"P{self.page}T{self.index}"[:31]


# --------------------------------------------------------------------------
# Pages helper: turn "1-3,5" into what camelot wants and a set we can track.
# --------------------------------------------------------------------------

def parse_pages(spec, total=None):
    """Return (camelot_pages_string, wanted_set_or_None). spec 'all' or '1,2,5' or '1-3'."""
    spec = (spec or "all").strip().lower()
    if spec in ("all", "", "*"):
        return "all", None
    wanted = []
    for token in spec.split(","):
        token = token.strip()
        if not token:
            continue
        if "-" in token:
            a, _, b = token.partition("-")
            try:
                lo, hi = int(a), int(b)
            except ValueError:
                raise SystemExit(f"Bad --pages value: {spec!r}")
            wanted.extend(range(lo, hi + 1))
        else:
            try:
                wanted.append(int(token))
            except ValueError:
                raise SystemExit(f"Bad --pages value: {spec!r}")
    wanted = sorted({p for p in wanted if p > 0})
    if not wanted:
        raise SystemExit(f"Bad --pages value: {spec!r}")
    if total is not None:
        wanted = [p for p in wanted if p <= total]
    return ",".join(map(str, wanted)), set(wanted)


def total_pages(pdf_path) -> int | None:
    try:
        import pypdfium2 as pdfium
        return len(pdfium.PdfDocument(str(pdf_path)))
    except Exception:
        return None


# --------------------------------------------------------------------------
# Extraction with camelot
# --------------------------------------------------------------------------

def _rows_from_df(df):
    df2 = df.fillna("").astype(str)
    return [[str(c).strip() for c in row] for row in df2.values.tolist()]


def _is_junk(rows):
    """A table needs at least 2 columns and some non-empty content."""
    cols = max((len(r) for r in rows), default=0)
    cells = sum(1 for r in rows for c in r if c)
    return cols < 2 or cells < 2


def extract_flavor(pdf_path, pages_str, flavor):
    """Run one camelot flavor. Returns list[(page:int|None, rows, accuracy|None)].

    Raises ImportError (engine missing) or Exception (per-run problem);
    callers catch and turn problems into warnings, never crashes.
    """
    import logging
    logging.getLogger("pdfminer").setLevel(logging.ERROR)  # hide parser chatter
    import camelot

    out = []
    tables = camelot.read_pdf(str(pdf_path), pages=pages_str, flavor=flavor,
                              suppress_stdout=True)
    for t in tables:
        report = getattr(t, "parsing_report", {}) or {}
        acc = report.get("accuracy")
        rows = _rows_from_df(t.df)
        if _is_junk(rows):
            continue
        # Stream mode on non-table pages produces confident-looking junk;
        # drop its weakest guesses instead of dumping noise on the buyer.
        if flavor == "stream" and acc is not None and acc < 40:
            continue
        page_no = report.get("page")
        out.append((int(page_no) if page_no is not None else None, rows,
                    round(float(acc), 1) if acc is not None else None))
    return out


def _friendly_extract_error(exc, flavor):
    text = str(exc).lower()
    if isinstance(exc, ImportError):
        return ("The PDF engine is not installed. Run install.bat first "
                "(or: pip install camelot-py==0.11.0 opencv-python-headless).")
    if "ghostscript" in text or "gs" == text.strip():
        return ("Ghostscript is missing (needed for the lined-grid mode). "
                "Install it free: winget install ArtifexSoftware.GhostScript "
                "- or rerun with --mode stream.")
    return f"{flavor} mode failed on this file: {exc}"


# --------------------------------------------------------------------------
# OCR fallback (optional flag; needs Tesseract + pypdfium2 + pytesseract)
# --------------------------------------------------------------------------

def extract_ocr_page(pdf_path, page_no):
    """Rasterize one page, OCR it, rebuild a rough table from word boxes.

    Returns rows (list of lists of strings) or None. Always reported LOW.
    """
    import pypdfium2 as pdfium
    import pytesseract

    doc = pdfium.PdfDocument(str(pdf_path))
    try:
        bitmap = doc[page_no - 1].render(scale=3)  # ~216 dpi, CPU-only
        image = bitmap.to_pil().convert("L")
    finally:
        doc.close()

    data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)

    words = []       # dicts: x0,x1,ycenter,h,text,line-key
    line_h = []
    n = len(data["text"])
    for k in range(n):
        txt = (data["text"][k] or "").strip()
        try:
            conf = float(data["conf"][k])
        except (TypeError, ValueError):
            conf = -1.0
        if not txt or conf < 0:
            continue
        left = int(data["left"][k]); top = int(data["top"][k])
        w = int(data["width"][k]); h = int(data["height"][k])
        key = (data["block_num"][k], data["par_num"][k], data["line_num"][k])
        words.append({"x0": left, "x1": left + w, "yc": top + h / 2.0,
                      "h": h, "text": txt, "key": key})
        line_h.append(h)
    if not words:
        return None

    med_h = statistics.median(line_h) or 12
    min_gap = max(8, int(med_h * 0.9))     # horizontal gap that means "new column"
    band_thr = max(4, int(med_h * 0.6))    # vertical distance that means "same row"

    # 1) group words into visual lines by tesseract's line keys
    lines = {}
    for w in words:
        lines.setdefault(w["key"], []).append(w)
    ordered_lines = sorted(lines.values(), key=lambda ws: min(w["yc"] for w in ws))

    # 2) merge close lines into row bands (OCR splits one row into 2+ lines often)
    bands = []
    for ws in ordered_lines:
        y = statistics.median(w["yc"] for w in ws)
        if bands and abs(y - bands[-1][0]) <= band_thr:
            bands[-1][1].extend(ws)
            bands[-1][0] = (bands[-1][0] + y) / 2.0
        else:
            bands.append([y, list(ws)])

    # 3) find column separators from gaps in word coverage across the whole page
    intervals = sorted((w["x0"], w["x1"]) for w in words)
    merged = [list(intervals[0])]
    for s, e in intervals[1:]:
        if s <= merged[-1][1] + 2:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    seps = [((merged[i][1] + merged[i + 1][0]) // 2)
            for i in range(len(merged) - 1)
            if merged[i + 1][0] - merged[i][1] >= min_gap]
    if not seps:
        return None  # single column of text - not a table, say nothing

    ncols = len(seps) + 1
    grid = {}
    for r_idx, (_, band_words) in enumerate(bands):
        for w in sorted(band_words, key=lambda w: w["x0"]):
            col = bisect.bisect_right(seps, (w["x0"] + w["x1"]) // 2)
            cell = grid.setdefault((r_idx, col), [])
            if not cell or cell[-1] != w["text"]:
                cell.append(w["text"])

    rows = [[" ".join(grid.get((r, c), [])) for c in range(ncols)]
            for r in range(len(bands))]
    rows = [r for r in rows if any(c for c in r)]  # drop fully blank rows
    if not rows:
        return None
    return rows


# --------------------------------------------------------------------------
# Orchestrator for one PDF
# --------------------------------------------------------------------------

def extract_pdf_full(pdf_path, pages_spec="all", mode="auto", ocr=False, log=print):
    """Extract tables from one PDF.

    Returns (results: list[TableResult], warnings: list[str]).
    Auto strategy: lined grid first; where it saw nothing, try spacing guess;
    where both saw nothing AND --ocr was passed, run OCR (graded LOW).
    """
    """Same contract as extract_pdf but stores method correctly end-to-end."""
    warnings = []
    tot = total_pages(pdf_path) if ocr else None
    pages_str, wanted = parse_pages(pages_spec, tot)
    by_page = {}

    def store(p, rows, acc, method):
        if wanted is not None and p not in wanted:
            return
        by_page.setdefault(p, []).append((rows, acc, method))

    if mode in ("auto", "lattice"):
        try:
            for p, rows, acc in extract_flavor(pdf_path, pages_str, "lattice"):
                store(p, rows, acc, "lattice")
        except Exception as exc:
            warnings.append(_friendly_extract_error(exc, "Lattice"))

    if mode in ("auto", "stream"):
        if mode == "auto":
            if wanted is not None:
                missing = sorted(wanted - set(by_page))
                targets = ",".join(map(str, missing)) if missing else None
            elif tot is not None:
                targets = ",".join(map(str, sorted(set(range(1, tot + 1)) - set(by_page)))) or None
            else:
                targets = "all"
        else:
            targets = pages_str
        if targets:
            try:
                for p, rows, acc in extract_flavor(pdf_path, targets, "stream"):
                    if p is not None and by_page.get(p):
                        continue
                    store(p, rows, acc, "stream")
            except Exception as exc:
                warnings.append(_friendly_extract_error(exc, "Stream"))

    if ocr:
        if wanted is None and tot is None:
            warnings.append("OCR skipped: could not count pages (need pypdfium2).")
        else:
            todo = sorted(wanted) if wanted is not None else list(range(1, tot + 1))
            for p in todo:
                if by_page.get(p):
                    continue
                try:
                    rows = extract_ocr_page(pdf_path, p)
                except ImportError:
                    warnings.append("OCR needs Tesseract installed + "
                                    "pip install pytesseract pypdfium2 pillow. See README.")
                    break
                except Exception as exc:
                    warnings.append(f"OCR failed on page {p}: {exc}")
                    continue
                if rows and not _is_junk(rows):
                    by_page.setdefault(p, []).append((rows, None, "ocr"))
                    log(f"  page {p}: OCR fallback used (graded LOW, always proofread)")

    results = []
    for p in sorted(by_page):
        for i, (rows, acc, method) in enumerate(by_page[p], start=1):
            results.append(TableResult(page=p, index=i, method=method,
                                       accuracy=acc, rows=rows))
    return results, warnings


# --------------------------------------------------------------------------
# Writers: Excel + preview.html
# --------------------------------------------------------------------------

def write_excel(results, xlsx_path, src_label, mode_label):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([APP])
    ws.append(["Source", src_label])
    ws.append(["Created", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append(["Mode", mode_label])
    ws.append([])
    header = ["Sheet", "PDF page", "Table #", "Method", "Confidence",
              "Accuracy %", "Rows x Cols", "Note"]
    ws.append(header)
    for c in range(1, len(header) + 1):
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)

    counts = {"HIGH": 0, "MEDIUM": 0, "LOW": 0}
    for t in results:
        counts[t.label] += 1
        ws.append([t.sheet_name, t.page, t.index, METHOD_LABELS[t.method],
                   t.label, t.accuracy if t.accuracy is not None else "-",
                   t.size, t.note])
    ws.column_dimensions["H"].width = 60

    widths = {"A": 10, "B": 9, "C": 8, "D": 20, "E": 12, "F": 11, "G": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    for t in results:
        sheet = wb.create_sheet(t.sheet_name)
        for r in t.rows:
            sheet.append(r)
        if t.rows:
            for c in range(1, len(t.rows[0]) + 1):
                sheet.cell(row=1, column=c).font = Font(bold=True)
        sheet.freeze_panes = "A2"
        longest = {}
        for r in t.rows[:200]:
            for i, v in enumerate(r, start=1):
                longest[i] = min(max(longest.get(i, 8), len(str(v)) + 2), 42)
        for i, w in longest.items():
            sheet.column_dimensions[sheet.cell(row=1, column=i).column_letter].width = w

    wb.save(xlsx_path)
    return counts


def write_preview(results, html_path, src_label, mode_label, xlsx_name):
    badge = {"HIGH": "#0a7d33", "MEDIUM": "#b26a00", "LOW": "#b00020"}
    parts = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        f"<title>{html.escape(APP)} - preview</title><style>",
        "body{font-family:Segoe UI,Arial,sans-serif;margin:24px;color:#222}",
        "h1{font-size:20px} h2{font-size:15px;margin-top:28px;border-bottom:1px solid #ddd;padding-bottom:4px}",
        ".badge{color:#fff;padding:2px 8px;border-radius:3px;font-size:12px;font-weight:bold}",
        ".note{color:#555;font-size:13px;margin:6px 0}", "table{border-collapse:collapse;font-size:13px}",
        "td,th{border:1px solid #bbb;padding:4px 10px}", "tr:first-child td{background:#f2f2f2;font-weight:600}",
        ".meta{color:#666;font-size:13px}", "</style></head><body>",
        f"<h1>{html.escape(APP)} - preview</h1>",
        f"<p class='meta'>Source: <b>{html.escape(src_label)}</b> &middot; Mode: {html.escape(mode_label)} "
        f"&middot; Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} &middot; "
        f"{len(results)} table(s) found</p>",
        f"<p class='meta'>Excel written next to this file as <b>{html.escape(xlsx_name)}</b>. "
        "This preview shows exactly what went into it.</p>",
    ]
    if not results:
        parts.append("<p style='color:#b00020'>No tables were found. Try another mode "
                     "(spacing guess), or the OCR flag for scanned pages.</p>")
    cur_page = None
    for t in results:
        if t.page != cur_page:
            cur_page = t.page
            parts.append(f"<h2>Page {t.page}</h2>")
        color = badge[t.label]
        acc = f"{t.accuracy:.1f}%" if t.accuracy is not None else "n/a"
        parts.append(
            f"<p style='margin-bottom:2px'><b>Table {t.index}</b> "
            f"&middot; {html.escape(METHOD_LABELS[t.method])} &middot; "
            f"<span class='badge' style='background:{color}'>{t.label}</span> "
            f"&middot; score {acc} &middot; {t.size}</p>"
            f"<p class='note'>{html.escape(t.note)}</p>"
            "<table>")
        for r_i, row in enumerate(t.rows):
            tag = "th" if r_i == 0 else "td"
            cells = "".join(f"<{tag}>{html.escape(c)}</{tag}>" for c in row)
            parts.append(f"<tr>{cells}</tr>")
        parts.append("</table>")
    parts.append("<p class='note' style='margin-top:30px'>Preview only - open the .xlsx "
                 "for the actual export. LOW-grade tables deserve a proofread.</p></body></html>")
    Path(html_path).write_text("\n".join(parts), encoding="utf-8")


# --------------------------------------------------------------------------
# Per-file pipeline shared by CLI and GUI
# --------------------------------------------------------------------------

MODE_CHOICES = ("auto", "lattice", "stream")


def process_pdf(src: Path, out_file: Path, *, pages="all", mode="auto",
                ocr=False, preview=True, preview_path=None, log=print):
    """Run the full pipeline for one PDF. Returns a stats dict; never raises
    for per-file problems - failures land in stats['error']."""
    stats = {"src": src, "ok": False, "xlsx": None, "preview": None,
             "tables": 0, "high": 0, "medium": 0, "low": 0,
             "warnings": [], "error": None}
    t0 = datetime.now()
    try:
        log(f"[run] {src.name}")
        results, warnings = extract_pdf_full(src, pages_spec=pages, mode=mode,
                                             ocr=ocr, log=log)
        stats["warnings"] = warnings
        for w in warnings:
            log(f"  [warn] {w}")

        out_file.parent.mkdir(parents=True, exist_ok=True)
        counts = write_excel(results, out_file, src.name, mode)
        stats.update(ok=True, xlsx=out_file, tables=len(results),
                     high=counts["HIGH"], medium=counts["MEDIUM"], low=counts["LOW"])

        if preview:
            pv = Path(preview_path) if preview_path else \
                out_file.with_name(out_file.stem + ".preview.html")
            write_preview(results, pv, src.name, mode, out_file.name)
            stats["preview"] = pv

        secs = (datetime.now() - t0).total_seconds()
        if results:
            log(f"  [done] {len(results)} table(s) -> {out_file.name} "
                f"(HIGH {counts['HIGH']} / MEDIUM {counts['MEDIUM']} / LOW {counts['LOW']}) "
                f"in {secs:.1f}s")
        else:
            log(f"  [done] no tables found ({secs:.1f}s). Try --mode stream, "
                "or --ocr for scanned pages.")
    except SystemExit:
        raise
    except Exception as exc:
        stats["error"] = str(exc)
        log(f"  [FAIL] {src.name}: {exc}")
    return stats


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def expand_inputs(inputs):
    """Files and/or folders -> list of PDF paths (folders scanned for *.pdf)."""
    pdfs = []
    for item in inputs:
        p = Path(item)
        if p.is_dir():
            found = sorted(p.glob("*.pdf"))
            pdfs.extend(found)
            if not found:
                print(f"[warn] no PDFs inside folder: {p}")
        elif p.is_file() and p.suffix.lower() == ".pdf":
            pdfs.append(p)
        else:
            print(f"[warn] skipping (not a PDF): {p}")
    # de-dup, keep order
    seen = set()
    unique = []
    for p in pdfs:
        rp = str(p.resolve()).lower()
        if rp not in seen:
            seen.add(rp)
            unique.append(p)
    return unique


def build_parser():
    ap = argparse.ArgumentParser(
        prog="pdf-tables-to-excel",
        description=f"{APP} v{VERSION} - pull tables out of PDFs into Excel. "
                    "Runs fully offline on your PC.")
    ap.add_argument("inputs", nargs="*", help="one or more PDF files and/or folders")
    ap.add_argument("-o", "--output", help="output .xlsx path (single PDF only)")
    ap.add_argument("--out-dir", help="folder for batch outputs (default: beside each PDF)")
    ap.add_argument("--pages", default="all",
                    help="pages to read: all, or like 1,3-5 (default: all)")
    ap.add_argument("--mode", choices=MODE_CHOICES, default="auto",
                    help="auto = lined grids first, spacing guess fills the rest "
                         "(default); lattice = lined grids only; stream = spacing guess only")
    ap.add_argument("--ocr", action="store_true",
                    help="OCR fallback for scanned pages that gave no tables "
                         "(slower; needs free Tesseract)")
    ap.add_argument("--no-preview", action="store_true",
                    help="skip the per-file preview.html")
    ap.add_argument("--gui", action="store_true", help="open the desktop window")
    ap.add_argument("--version", action="version", version=f"{APP} {VERSION}")
    return ap


def run_cli(args):
    pdfs = expand_inputs(args.inputs)
    if not pdfs:
        build_parser().print_help()
        return 1
    if args.output and len(pdfs) > 1:
        print("[error] -o/--output works with exactly one PDF. "
              "For many files use --out-dir.")
        return 1

    exit_code = 0
    totals = {"files_ok": 0, "files_fail": 0, "tables": 0}
    for i, src in enumerate(pdfs, start=1):
        if args.output:
            out_file = Path(args.output)
        else:
            base = Path(args.out_dir) if args.out_dir else src.parent
            out_file = base / (src.stem + ".xlsx")
            guard = 1
            while out_file.exists() and args.out_dir:
                out_file = base / f"{src.stem}_{guard}.xlsx"
                guard += 1
        if len(pdfs) > 1:
            print(f"\n=== file {i}/{len(pdfs)} ===")
        stats = process_pdf(src, out_file, pages=args.pages, mode=args.mode,
                            ocr=args.ocr, preview=not args.no_preview)
        if stats["ok"]:
            totals["files_ok"] += 1
            totals["tables"] += stats["tables"]
        else:
            totals["files_fail"] += 1

    print(f"\nSummary: {totals['files_ok']} file(s) converted, "
          f"{totals['tables']} table(s) total, {totals['files_fail']} failed.")
    if totals["files_fail"]:
        exit_code = 2 if totals["files_ok"] else 1
    return exit_code


# --------------------------------------------------------------------------
# GUI (tkinter, standard library)
# --------------------------------------------------------------------------

def launch_gui():
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, scrolledtext, ttk
    except ImportError:
        print("tkinter is missing. Reinstall Python with the tcl/tk option ticked, "
              "or use the command line instead (python tables.py --help).")
        return 1

    MODE_DISPLAY = {
        "Auto (recommended): lined grids first, then spacing": "auto",
        "Lined grids only (most accurate)": "lattice",
        "Spacing guess only (no Ghostscript needed)": "stream",
    }
    DISPLAY_MODE = {v: k for k, v in MODE_DISPLAY.items()}

    app = tk.Tk()
    app.title(f"{APP} v{VERSION}")
    app.minsize(720, 520)

    files_var = tk.Variable(value=[])
    q = queue.Queue()
    running = {"flag": False}

    main = ttk.Frame(app, padding=10)
    main.pack(fill="both", expand=True)

    # file list
    top = ttk.Frame(main)
    top.pack(fill="x")
    ttk.Label(top, text="PDFs to convert:").pack(anchor="w")

    list_frame = ttk.Frame(main)
    list_frame.pack(fill="both", expand=True, pady=(4, 8))
    lb = tk.Listbox(list_frame, selectmode="extended", listvariable=files_var,
                    height=7)
    sb = ttk.Scrollbar(list_frame, orient="vertical", command=lb.yview)
    lb.config(yscrollcommand=sb.set)
    lb.pack(side="left", fill="both", expand=True)
    sb.pack(side="right", fill="y")

    btns = ttk.Frame(main)
    btns.pack(fill="x", pady=(0, 8))

    def add_files():
        picks = filedialog.askopenfilenames(title="Pick PDF files",
                                            filetypes=[("PDF files", "*.pdf")])
        if picks:
            files_var.set(list(files_var.get()) + [Path(p).__str__() for p in picks])

    def add_folder():
        d = filedialog.askdirectory(title="Pick a folder of PDFs")
        if d:
            files_var.set(list(files_var.get()) + [str(Path(d))])

    def remove_selected():
        sel = [lb.index(i) for i in lb.curselection()]
        keep = [f for i, f in enumerate(files_var.get()) if i not in sel]
        files_var.set(keep)

    ttk.Button(btns, text="Add PDFs...", command=add_files).pack(side="left")
    ttk.Button(btns, text="Add folder...", command=add_folder).pack(side="left", padx=6)
    ttk.Button(btns, text="Remove selected", command=remove_selected).pack(side="left")
    ttk.Button(btns, text="Clear list", command=lambda: files_var.set([])).pack(side="left", padx=6)

    # options
    opts = ttk.LabelFrame(main, text="Options", padding=8)
    opts.pack(fill="x", pady=(0, 8))

    out_dir_var = tk.StringVar()
    mode_var = tk.StringVar(value=list(MODE_DISPLAY)[0])
    pages_var = tk.StringVar(value="all")
    ocr_var = tk.BooleanVar(value=False)
    preview_var = tk.BooleanVar(value=True)

    ttk.Label(opts, text="Save spreadsheets to (blank = same folder as each PDF):")\
        .grid(row=0, column=0, sticky="w")
    ttk.Entry(opts, textvariable=out_dir_var, width=44).grid(row=1, column=0, sticky="we", padx=(0, 6))

    def browse_out():
        d = filedialog.askdirectory(title="Output folder")
        if d:
            out_dir_var.set(d)

    ttk.Button(opts, text="Browse...", command=browse_out).grid(row=1, column=1)
    ttk.Label(opts, text="How to read tables:").grid(row=0, column=2, padx=(16, 4), sticky="e")
    ttk.OptionMenu(opts, mode_var, mode_var.get(), *MODE_DISPLAY.keys())\
        .grid(row=1, column=2, padx=(16, 4))
    ttk.Label(opts, text="Pages:").grid(row=0, column=3, padx=(16, 4), sticky="e")
    ttk.Entry(opts, textvariable=pages_var, width=10).grid(row=1, column=3, padx=(16, 0))
    ttk.Checkbutton(opts, text="Try OCR on scanned pages (slower; needs free Tesseract)",
                    variable=ocr_var).grid(row=2, column=0, columnspan=3, sticky="w", pady=(6, 0))
    ttk.Checkbutton(opts, text="Write a preview.html so I can check each table first",
                    variable=preview_var).grid(row=3, column=0, columnspan=3, sticky="w")

    opts.columnconfigure(0, weight=1)

    # run + log
    run_row = ttk.Frame(main)
    run_row.pack(fill="x")
    run_btn = ttk.Button(run_row, text="Convert")
    run_btn.pack(side="left")
    open_btn = ttk.Button(run_row, text="Open output folder", state="disabled")
    open_btn.pack(side="left", padx=8)
    status_lbl = ttk.Label(run_row, text="")
    status_lbl.pack(side="left", padx=8)

    log_box = scrolledtext.ScrolledText(main, height=12, state="disabled",
                                        font=("Consolas", 9))
    log_box.pack(fill="both", expand=True, pady=(8, 0))

    def log(msg):
        q.put(msg)

    def poll():
        drained = False
        while True:
            try:
                msg = q.get_nowait()
            except queue.Empty:
                break
            drained = True
            if msg == "__DONE__":
                finish()
                continue
            log_box.config(state="normal")
            log_box.insert("end", msg + "\n")
            log_box.see("end")
            log_box.config(state="disabled")
        if running["flag"]:
            app.after(120, poll)

    last_outdir = {"path": None}

    def finish():
        running["flag"] = False
        run_btn.config(state="normal")
        status_lbl.config(text="Done.")
        try:
            log_box.config(state="normal")
            log_box.insert("end", "\nAll finished.\n")
            log_box.see("end")
            log_box.config(state="disabled")
        finally:
            pass
        if last_outdir["path"]:
            open_btn.config(state="normal")

    def worker(items, out_dir, mode, pages, ocr, preview):
        ok = fail = tabs = 0
        for i, item in enumerate(items, 1):
            p = Path(item)
            sources = sorted(p.glob("*.pdf")) if p.is_dir() else [p]
            for src in sources:
                log(f"\n=== {i}/{len(items)}: {src.name} ===")
                out_file = (Path(out_dir) if out_dir else src.parent) / (src.stem + ".xlsx")
                stats = process_pdf(src, out_file, pages=pages, mode=mode,
                                    ocr=ocr, preview=preview, log=log)
                if stats["ok"]:
                    ok += 1
                    tabs += stats["tables"]
                    last_outdir["path"] = str(out_file.parent)
                else:
                    fail += 1
        log(f"\nSummary: {ok} converted, {tabs} table(s), {fail} failed.")
        q.put("__DONE__")

    def on_run():
        items = files_var.get()
        if not items:
            messagebox.showinfo(APP, "Add at least one PDF or folder first.")
            return
        if ocr_var.get():
            confirm = messagebox.askyesno(
                APP, "OCR mode reads scanned pages letter by letter. It is slow.\n"
                     "Continue?")
            if not confirm:
                return
        running["flag"] = True
        run_btn.config(state="disabled")
        open_btn.config(state="disabled")
        status_lbl.config(text="Working...")
        mode = MODE_DISPLAY[mode_var.get()]
        threading.Thread(target=worker, daemon=True, args=(
            items, out_dir_var.get().strip(), mode,
            pages_var.get().strip() or "all",
            ocr_var.get(), preview_var.get())).start()
        app.after(120, poll)

    def on_open():
        if last_outdir["path"] and os.path.isdir(last_outdir["path"]):
            try:
                os.startfile(last_outdir["path"])  # Windows
            except AttributeError:
                import subprocess
                subprocess.Popen(["xdg-open", last_outdir["path"]])

    run_btn.config(command=on_run)
    open_btn.config(command=on_open)

    log_box.config(state="normal")
    log_box.insert("end", f"{APP} v{VERSION}\n"
                           "Everything runs on this PC. Your files never leave it.\n"
                           "1) Add PDFs or a folder  2) Convert  3) Check the preview\n")
    log_box.config(state="disabled")

    app.mainloop()
    return 0


# --------------------------------------------------------------------------

def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv or "--gui" in argv:
        argv = [a for a in argv if a != "--gui"]
        if argv:  # --gui combined with inputs is ambiguous; refuse politely
            print("[error] Use either file names OR --gui, not both.")
            return 1
        return launch_gui()
    args = build_parser().parse_args(argv)
    return run_cli(args)


if __name__ == "__main__":
    sys.exit(main())
